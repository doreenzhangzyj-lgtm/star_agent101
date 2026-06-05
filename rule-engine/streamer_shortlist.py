#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
潜力主播海选系统 — 单文件可运行版

依赖: pip install openpyxl

用法:
  python3 streamer_shortlist.py
  python3 streamer_shortlist.py -i 指标.csv -a ASR目录 -p 潜能.xlsx -o 结果.json
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None  # type: ignore

# =============================================================================
# 配置
# =============================================================================

SYSTEM_NAME = "潜力主播海选系统"
VERSION = "2.0.0"

DEFAULT_CSV = "/Users/bytedance/Desktop/黑客/Agent大星探底池测试第二次.csv"
DEFAULT_ASR_DIR = "/Users/bytedance/Desktop/黑客/ASR目录"
DEFAULT_EXCEL = "/Users/bytedance/Desktop/黑客/主播潜能判断表.xlsx"
DEFAULT_OUTPUT = "/Users/bytedance/Desktop/黑客/shortlist_result.json"

MODULE_MAX: Dict[str, float] = {
    "user_consumption": 60.0,
    "business_capability": 20.0,
    "content_stability": 10.0,
    "streamer_potential": 10.0,
}

CONSUMPTION_FIELD_WEIGHTS: Dict[str, float] = {
    "avg_online": 0.18,
    "max_online": 0.12,
    "watch_count": 0.15,
    "avg_session_watch_min": 0.10,
    "avg_user_watch_min": 0.18,
    "like_rate": 0.10,
    "comment_rate": 0.09,
    "like_engagement": 0.04,
    "comment_engagement": 0.04,
}

POTENTIAL_FIELD_WEIGHTS: Dict[str, float] = {
    "cooperation": 0.20,
    "persona_clarity": 0.20,
    "style_stability": 0.20,
    "values_positive": 0.20,
    "expression": 0.20,
}

LEADS_TOP_PERCENTILE = 0.80
STABILITY_MIN_SESSIONS_PER_WEEK = 2
STABILITY_MAX_AVG_ONLINE_SPREAD = 100.0
SWEAR_WORD_MAX_ALLOWED = 2
PASS_TOTAL_SCORE = 60.0
INCLUDE_REJECTED_IN_OUTPUT = True
POTENTIAL_SCORE_MAX = 5.0

GUIDANCE_KEYWORDS: Tuple[str, ...] = (
    "步骤", "方法", "教程", "怎么做", "如何做", "操作指南", "教学",
    "手把手", "配置", "设置", "技巧", "攻略", "流程", "指南", "演示",
    "how to", "step by step", "tutorial", "guide", "teach", "tips",
    "method", "workflow", "setup", "configure", "instruction", "demo",
    "walkthrough", "show you", "let me show",
)

PROFANITY_WORDS: Tuple[str, ...] = (
    "他妈", "尼玛", "你妈", "傻逼", "傻B", "操你", "草你", "滚蛋",
    "去死", "贱人", "狗屎", "废物", "妈的", "TMD", "tmd", "NMD", "nmd",
    "fuck", "fucking", "shit", "bitch", "asshole", "damn",
)

ASR_EXTENSIONS = (".txt", ".asr", ".text", ".md")

CSV_COLUMN_ALIASES: Dict[str, str] = {
    "编号": "serial_no", "serial_no": "serial_no", "case_no": "serial_no",
    "p_date": "p_date", "date": "p_date",
    "主播id": "author_id", "author_id": "author_id",
    "主播昵称": "author_nickname", "nickname": "author_nickname",
    "直播间id": "live_room_id", "live_room_id": "live_room_id",
    "整场平均在线人数": "avg_online", "整场最大在线人数": "max_online",
    "整场最大在线人": "max_online", "直播时观看次数": "watch_count",
    "次均看播时长：分钟": "avg_session_watch_min",
    "次均看播时长:分钟": "avg_session_watch_min", "次均看播时长": "avg_session_watch_min",
    "人均看播时长：分钟": "avg_user_watch_min",
    "人均看播时长:分钟": "avg_user_watch_min", "人均看播时长": "avg_user_watch_min",
    "点赞率": "like_rate", "评论率": "comment_rate", "总线索数": "total_leads",
    "点赞pv": "like_pv", "点赞uv": "like_uv", "评论pv": "comment_pv", "评论uv": "comment_uv",
}

CSV_REQUIRED_COLUMNS = (
    "serial_no", "p_date", "author_id", "live_room_id", "avg_online", "max_online",
    "watch_count", "avg_session_watch_min", "avg_user_watch_min",
    "like_rate", "comment_rate", "total_leads",
)

EXCEL_COLUMN_ALIASES: Dict[str, str] = {
    "编号": "serial_no", "serial_no": "serial_no",
    "主播id": "author_id", "author_id": "author_id",
    "配合度": "cooperation", "人设清晰度": "persona_clarity",
    "内容风格稳定性": "style_stability",
    "价值观风险": "values_positive", "价值观正向度": "values_positive",
    "表达能力": "expression",
}

EXCEL_REQUIRED_COLUMNS = (
    "serial_no", "cooperation", "persona_clarity", "style_stability",
    "values_positive", "expression",
)

LEVEL_SCORE_MAP: Dict[str, float] = {
    "高": 100.0, "中": 60.0, "低": 20.0,
    "high": 100.0, "medium": 60.0, "mid": 60.0, "low": 20.0,
}


# =============================================================================
# 数据模型
# =============================================================================

@dataclass
class LiveMetricsRow:
    serial_no: str
    p_date: str
    author_id: str
    author_nickname: str
    live_room_id: str
    avg_online: float
    max_online: float
    watch_count: float
    avg_session_watch_min: float
    avg_user_watch_min: float
    like_rate: float
    comment_rate: float
    like_pv: float
    like_uv: float
    comment_pv: float
    comment_uv: float
    total_leads: float
    row_index: int = 0

    @property
    def like_pv_uv_ratio(self) -> float:
        return self.like_pv / self.like_uv if self.like_uv > 0 else 0.0

    @property
    def comment_pv_uv_ratio(self) -> float:
        return self.comment_pv / self.comment_uv if self.comment_uv > 0 else 0.0


@dataclass
class PotentialProfile:
    serial_no: str
    author_id: str
    cooperation: float
    persona_clarity: float
    style_stability: float
    values_risk: float
    expression: float


@dataclass
class AnchorSession:
    serial_no: str
    author_id: str
    live_room_id: str
    author_nickname: str
    representative: LiveMetricsRow
    week_rows: List[LiveMetricsRow] = field(default_factory=list)
    sample_asr: str = ""
    potential: Optional[PotentialProfile] = None


@dataclass
class HardRejectResult:
    rejected: bool
    reasons: List[str] = field(default_factory=list)
    profanity_count: int = 0
    guidance_info_count: int = 0


@dataclass
class ModuleScores:
    user_consumption: float = 0.0
    business_capability: float = 0.0
    content_stability: float = 0.0
    streamer_potential: float = 0.0
    detail: Dict[str, Any] = field(default_factory=dict)

    @property
    def total(self) -> float:
        return round(
            self.user_consumption + self.business_capability
            + self.content_stability + self.streamer_potential, 2,
        )

    def as_dict(self) -> Dict[str, Any]:
        return {
            "user_consumption": round(self.user_consumption, 2),
            "business_capability": round(self.business_capability, 2),
            "content_stability": round(self.content_stability, 2),
            "streamer_potential": round(self.streamer_potential, 2),
            "total": self.total,
            "detail": self.detail,
        }


@dataclass
class ScoredAnchor:
    session: AnchorSession
    hard_reject: HardRejectResult
    modules: ModuleScores
    status: str

    def is_passed(self) -> bool:
        return self.status == "入围"


@dataclass
class ShortlistResult:
    input_csv: str
    asr_dir: str
    potential_excel: str
    items: List[ScoredAnchor] = field(default_factory=list)
    total_input: int = 0
    total_passed: int = 0


# =============================================================================
# 工具函数
# =============================================================================

def _clamp(v: float, lo: float = 0.0, hi: float = 100.0) -> float:
    return max(lo, min(hi, v))


def _parse_float(raw: str, default: float = 0.0) -> float:
    if raw is None:
        return default
    text = str(raw).strip().replace(",", "").replace("%", "")
    if not text or text in ("-", "—", "NA", "null", "None"):
        return default
    try:
        return float(text)
    except ValueError:
        m = re.search(r"-?\d+\.?\d*", text)
        return float(m.group()) if m else default


def _parse_pv_uv_cell(raw: str) -> Tuple[float, float]:
    text = str(raw).strip()
    if not text:
        return 0.0, 0.0
    for sep in ("/", "|", "／"):
        if sep in text:
            parts = [p.strip() for p in text.split(sep, 1)]
            if len(parts) == 2:
                return _parse_float(parts[0]), _parse_float(parts[1])
    return _parse_float(text), 0.0


def _norm_csv_header(name: str) -> str:
    key = name.strip().lower().replace(" ", "").replace("：", ":")
    if "点赞pv" in key and "uv" in key:
        return "like_pv_uv_combo"
    if "评论pv" in key and "uv" in key:
        return "comment_pv_uv_combo"
    return CSV_COLUMN_ALIASES.get(key, key)


def _decode_csv_bytes(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _serial_sort_key(serial_no: str) -> int:
    try:
        return int(str(serial_no).strip())
    except ValueError:
        return 0


def _serial_variants(serial_no: str) -> List[str]:
    s = str(serial_no).strip()
    variants = [s]
    try:
        n = int(s)
        variants.extend([str(n), f"{n:02d}"])
    except ValueError:
        pass
    return list(dict.fromkeys(variants))


def _read_asr_text(path: Path) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin-1"):
        try:
            return path.read_text(encoding=enc).strip()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_bytes().decode("utf-8", errors="replace").strip()


def _norm_excel_header(cell: object) -> str:
    return EXCEL_COLUMN_ALIASES.get(str(cell or "").strip().lower().replace(" ", ""), "")


def _norm_serial(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _parse_excel_score(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        raw = float(value)
        return raw / POTENTIAL_SCORE_MAX * 100.0 if raw <= 5 else min(100.0, raw)
    text = str(value).strip()
    if text in LEVEL_SCORE_MAP:
        return LEVEL_SCORE_MAP[text]
    try:
        raw = float(text)
        return raw / POTENTIAL_SCORE_MAX * 100.0 if raw <= 5 else min(100.0, raw)
    except ValueError:
        return 0.0


def _parse_date(s: str) -> Optional[datetime]:
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime((s or "").strip(), fmt)
        except ValueError:
            continue
    return None


def _iso_week_key(dt: datetime) -> str:
    y, w, _ = dt.isocalendar()
    return f"{y}-W{w:02d}"


# =============================================================================
# CSV 解析
# =============================================================================

def _read_csv_rows(path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    reader = csv.DictReader(io.StringIO(_decode_csv_bytes(path), newline=""))
    if not reader.fieldnames:
        raise ValueError("CSV 无表头")
    headers = [_norm_csv_header(h) for h in reader.fieldnames]
    rows = []
    for row in reader:
        rows.append({
            std: (row.get(orig) or "").strip()
            for orig, std in zip(reader.fieldnames, headers)
        })
    return headers, rows


def _validate_csv_headers(headers: Iterable[str]) -> None:
    missing = [c for c in CSV_REQUIRED_COLUMNS if c not in set(headers)]
    if missing:
        raise ValueError(f"CSV 缺少列: {', '.join(missing)}")


def _row_to_metrics(row: Dict[str, str], idx: int) -> Optional[LiveMetricsRow]:
    serial_no = str(row.get("serial_no", "")).strip()
    author_id = row.get("author_id", "")
    if not serial_no or not author_id:
        return None

    like_pv, like_uv = _parse_pv_uv_cell(row.get("like_pv_uv_combo", ""))
    if row.get("like_pv") or row.get("like_uv"):
        like_pv = _parse_float(row.get("like_pv", like_pv))
        like_uv = _parse_float(row.get("like_uv", like_uv))

    comment_pv, comment_uv = _parse_pv_uv_cell(row.get("comment_pv_uv_combo", ""))
    if row.get("comment_pv") or row.get("comment_uv"):
        comment_pv = _parse_float(row.get("comment_pv", comment_pv))
        comment_uv = _parse_float(row.get("comment_uv", comment_uv))

    return LiveMetricsRow(
        serial_no=serial_no, p_date=row.get("p_date", ""), author_id=author_id,
        author_nickname=row.get("author_nickname", ""), live_room_id=row.get("live_room_id", ""),
        avg_online=_parse_float(row.get("avg_online", "")),
        max_online=_parse_float(row.get("max_online", "")),
        watch_count=_parse_float(row.get("watch_count", "")),
        avg_session_watch_min=_parse_float(row.get("avg_session_watch_min", "")),
        avg_user_watch_min=_parse_float(row.get("avg_user_watch_min", "")),
        like_rate=_parse_float(row.get("like_rate", "")),
        comment_rate=_parse_float(row.get("comment_rate", "")),
        like_pv=like_pv, like_uv=like_uv, comment_pv=comment_pv, comment_uv=comment_uv,
        total_leads=_parse_float(row.get("total_leads", "")),
        row_index=idx,
    )


def parse_metrics_csv(path: str | Path) -> List[LiveMetricsRow]:
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"CSV 不存在: {file_path}")
    headers, rows = _read_csv_rows(file_path)
    _validate_csv_headers(headers)
    metrics = []
    for idx, row in enumerate(rows, start=2):
        m = _row_to_metrics(row, idx)
        if m:
            metrics.append(m)
    if not metrics:
        raise ValueError("CSV 无有效数据")
    return metrics


def build_anchor_sessions(rows: List[LiveMetricsRow]) -> List[AnchorSession]:
    by_author: Dict[str, List[LiveMetricsRow]] = defaultdict(list)
    for r in rows:
        by_author[r.author_id].append(r)
    sessions = [
        AnchorSession(
            serial_no=r.serial_no, author_id=r.author_id, live_room_id=r.live_room_id,
            author_nickname=r.author_nickname, representative=r, week_rows=by_author[r.author_id],
        )
        for r in rows
    ]
    return sorted(sessions, key=lambda s: (_serial_sort_key(s.serial_no), s.live_room_id))


# =============================================================================
# ASR 加载
# =============================================================================

def load_asr_directory(asr_dir: str | Path) -> Dict[str, str]:
    root = Path(asr_dir)
    if not root.is_dir():
        raise FileNotFoundError(f"ASR 目录不存在: {root}")

    index: Dict[str, str] = {}

    def register(key: str, text: str) -> None:
        if key and text:
            index[key] = text

    for entry in root.iterdir():
        name = entry.name.strip()
        if entry.is_dir():
            inner = entry / f"{name}.txt"
            if inner.is_file():
                register(name, _read_asr_text(inner))
            else:
                for ext in ASR_EXTENSIONS:
                    for fp in entry.glob(f"*{ext}"):
                        register(name, _read_asr_text(fp))
                        break
        elif entry.suffix.lower() in ASR_EXTENSIONS:
            register(entry.stem, _read_asr_text(entry))

    return index


def resolve_asr(index: Dict[str, str], serial_no: str, live_room_id: str) -> str:
    room = str(live_room_id).strip()
    for sn in _serial_variants(serial_no):
        key = f"{sn}_{room}"
        if key in index:
            return index[key]
    return ""


# =============================================================================
# Excel 解析
# =============================================================================

def parse_potential_excel(path: str | Path) -> Dict[str, PotentialProfile]:
    if load_workbook is None:
        raise ImportError("请先安装: pip install openpyxl")
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"Excel 不存在: {file_path}")

    wb = load_workbook(file_path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("Excel 为空")

    headers = [_norm_excel_header(h) for h in rows[0]]
    present = set(headers)
    missing = [c for c in EXCEL_REQUIRED_COLUMNS if c not in present]
    if missing:
        raise ValueError(f"Excel 缺少列: {', '.join(missing)}")

    col_idx = {name: headers.index(name) for name in EXCEL_REQUIRED_COLUMNS}
    profiles: Dict[str, PotentialProfile] = {}
    for row in rows[1:]:
        if not row:
            continue
        serial_no = _norm_serial(row[col_idx["serial_no"]])
        if not serial_no:
            continue
        author_id = ""
        if "author_id" in headers:
            author_id = str(row[headers.index("author_id")] or "").strip()
        profiles[serial_no] = PotentialProfile(
            serial_no=serial_no, author_id=author_id,
            cooperation=_parse_excel_score(row[col_idx["cooperation"]]),
            persona_clarity=_parse_excel_score(row[col_idx["persona_clarity"]]),
            style_stability=_parse_excel_score(row[col_idx["style_stability"]]),
            values_risk=_parse_excel_score(row[col_idx["values_positive"]]),
            expression=_parse_excel_score(row[col_idx["expression"]]),
        )
    return profiles


def get_profile(profiles: Dict[str, PotentialProfile], serial_no: str) -> Optional[PotentialProfile]:
    key = _norm_serial(serial_no)
    if key in profiles:
        return profiles[key]
    try:
        return profiles.get(str(int(key)))
    except ValueError:
        return None


# =============================================================================
# 打分
# =============================================================================

def count_profanity(text: str) -> int:
    if not text:
        return 0
    total, lower = 0, text.lower()
    for word in PROFANITY_WORDS:
        w, start = word.lower(), 0
        while True:
            pos = lower.find(w, start)
            if pos < 0:
                break
            total += 1
            start = pos + len(w)
    return total


def count_guidance_info(text: str) -> int:
    if not text:
        return 0
    lower = text.lower()
    return sum(1 for kw in GUIDANCE_KEYWORDS if kw in text or kw.lower() in lower)


def evaluate_hard_reject(asr: str) -> HardRejectResult:
    profanity, guidance = count_profanity(asr), count_guidance_info(asr)
    reasons = []
    if profanity > SWEAR_WORD_MAX_ALLOWED:
        reasons.append(f"脏话次数 {profanity} 超过上限 {SWEAR_WORD_MAX_ALLOWED}")
    if guidance == 0:
        reasons.append("指导方法信息为 0")
    return HardRejectResult(bool(reasons), reasons, profanity, guidance)


def _percentile_rank(values: List[float], value: float) -> float:
    if not values:
        return 0.0
    if len(values) == 1:
        return 100.0 if value >= values[0] else 0.0
    sv = sorted(values)
    below = sum(1 for v in sv if v < value)
    equal = sum(1 for v in sv if v == value)
    return _clamp((below + 0.5 * equal) / len(sv) * 100.0)


def _norm_potential(raw: float) -> float:
    if raw <= 0:
        return 0.0
    return _clamp(raw) if raw > 10 else _clamp(raw / POTENTIAL_SCORE_MAX * 100.0)


def score_user_consumption(rep: LiveMetricsRow, cohort: List[LiveMetricsRow]) -> Tuple[float, Dict[str, float]]:
    fields = {
        "avg_online": [r.avg_online for r in cohort],
        "max_online": [r.max_online for r in cohort],
        "watch_count": [r.watch_count for r in cohort],
        "avg_session_watch_min": [r.avg_session_watch_min for r in cohort],
        "avg_user_watch_min": [r.avg_user_watch_min for r in cohort],
        "like_rate": [r.like_rate for r in cohort],
        "comment_rate": [r.comment_rate for r in cohort],
        "like_engagement": [r.like_pv_uv_ratio for r in cohort],
        "comment_engagement": [r.comment_pv_uv_ratio for r in cohort],
    }
    sample = {
        "avg_online": rep.avg_online, "max_online": rep.max_online,
        "watch_count": rep.watch_count, "avg_session_watch_min": rep.avg_session_watch_min,
        "avg_user_watch_min": rep.avg_user_watch_min, "like_rate": rep.like_rate,
        "comment_rate": rep.comment_rate, "like_engagement": rep.like_pv_uv_ratio,
        "comment_engagement": rep.comment_pv_uv_ratio,
    }
    sub, wsum = {}, 0.0
    for field, weight in CONSUMPTION_FIELD_WEIGHTS.items():
        pct = _percentile_rank(fields[field], sample[field])
        sub[field] = round(pct, 2)
        wsum += pct * weight
    return round(wsum / 100.0 * MODULE_MAX["user_consumption"], 2), sub


def score_business_capability(rep: LiveMetricsRow, cohort: List[LiveMetricsRow]) -> Tuple[float, Dict[str, float]]:
    pct = _percentile_rank([r.total_leads for r in cohort], rep.total_leads)
    threshold = (1.0 - LEADS_TOP_PERCENTILE) * 100.0
    score = MODULE_MAX["business_capability"] if pct >= threshold else (
        (pct / threshold) * MODULE_MAX["business_capability"] if threshold else 0.0
    )
    return round(score, 2), {
        "total_leads": rep.total_leads, "leads_percentile": round(pct, 2),
        "top_80_percentile_threshold": round(threshold, 2),
    }


def score_content_stability(session: AnchorSession) -> Tuple[float, Dict[str, object]]:
    by_week: Dict[str, List[LiveMetricsRow]] = defaultdict(list)
    for row in session.week_rows:
        dt = _parse_date(row.p_date)
        if dt:
            by_week[_iso_week_key(dt)].append(row)

    best_ratio, best_detail = 0.0, {"passed": False, "weeks_evaluated": len(by_week)}
    for week_key, week_rows in by_week.items():
        session_count = len({r.p_date for r in week_rows})
        avgs = [r.avg_online for r in week_rows]
        spread = max(avgs) - min(avgs) if avgs else 0.0
        ok1 = session_count >= STABILITY_MIN_SESSIONS_PER_WEEK
        ok2 = spread <= STABILITY_MAX_AVG_ONLINE_SPREAD
        ratio = (0.5 if ok1 else 0.0) + (0.5 if ok2 else 0.0)
        if ratio > best_ratio:
            best_ratio = ratio
            best_detail = {
                "passed": ok1 and ok2, "week": week_key, "session_count": session_count,
                "avg_online_spread": round(spread, 2),
                "session_requirement_met": ok1, "spread_requirement_met": ok2,
            }
    return round(best_ratio * MODULE_MAX["content_stability"], 2), best_detail


def score_streamer_potential(profile: Optional[PotentialProfile]) -> Tuple[float, Dict[str, float]]:
    if not profile:
        return 0.0, {"missing_excel_profile": True}
    raw = {
        "cooperation": _norm_potential(profile.cooperation),
        "persona_clarity": _norm_potential(profile.persona_clarity),
        "style_stability": _norm_potential(profile.style_stability),
        "values_positive": _clamp(profile.values_risk),
        "expression": _norm_potential(profile.expression),
    }
    weighted = sum(raw[k] * POTENTIAL_FIELD_WEIGHTS[k] for k in POTENTIAL_FIELD_WEIGHTS)
    return round(weighted / 100.0 * MODULE_MAX["streamer_potential"], 2), raw


def score_anchor(session: AnchorSession, cohort: List[LiveMetricsRow]) -> ScoredAnchor:
    hard = evaluate_hard_reject(session.sample_asr)
    if hard.rejected:
        return ScoredAnchor(session, hard, ModuleScores(), "硬性淘汰")

    c, cd = score_user_consumption(session.representative, cohort)
    b, bd = score_business_capability(session.representative, cohort)
    s, sd = score_content_stability(session)
    p, pd = score_streamer_potential(session.potential)
    modules = ModuleScores(c, b, s, p, {
        "user_consumption": cd, "business_capability": bd,
        "content_stability": sd, "streamer_potential": pd,
    })
    status = "入围" if modules.total >= PASS_TOTAL_SCORE else "综合淘汰"
    return ScoredAnchor(session, hard, modules, status)


def run_scoring(sessions: List[AnchorSession], input_csv: str, asr_dir: str, excel_path: str) -> ShortlistResult:
    cohort = [s.representative for s in sessions]
    items = [score_anchor(s, cohort) for s in sessions]
    return ShortlistResult(
        input_csv, asr_dir, excel_path, items, len(sessions),
        sum(1 for x in items if x.is_passed()),
    )


# =============================================================================
# JSON 导出
# =============================================================================

def _item_to_dict(item: ScoredAnchor) -> Dict[str, Any]:
    s, r = item.session, item.session.representative
    return {
        "编号": s.serial_no,
        "主播ID": s.author_id,
        "直播间ID": s.live_room_id,
        "asr_key": f"{s.serial_no}_{s.live_room_id}",
        "主播昵称": s.author_nickname,
        "p_date": r.p_date,
        "整场平均在线人数": r.avg_online,
        "人均看播时长：分钟": r.avg_user_watch_min,
        "点赞率": r.like_rate,
        "评论率": r.comment_rate,
        "sample_asr": s.sample_asr,
        "status": item.status,
        "hard_reject": {
            "rejected": item.hard_reject.rejected,
            "reasons": item.hard_reject.reasons,
            "profanity_count": item.hard_reject.profanity_count,
            "guidance_info_count": item.hard_reject.guidance_info_count,
        },
        "scores": item.modules.as_dict(),
    }


def export_json(result: ShortlistResult, output_path: str | Path) -> Path:
    items = result.items if INCLUDE_REJECTED_IN_OUTPUT else [i for i in result.items if i.is_passed()]
    payload = {
        "meta": {
            "system": SYSTEM_NAME, "version": VERSION,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "input_csv": result.input_csv, "asr_dir": result.asr_dir,
            "potential_excel": result.potential_excel,
            "total_input": result.total_input, "total_passed": result.total_passed,
            "total_output": len(items),
        },
        "results": [_item_to_dict(i) for i in items],
    }
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return out


# =============================================================================
# 主流程
# =============================================================================

def run_pipeline(csv_path: Path, asr_dir: Path, excel_path: Path, output_path: Path) -> ShortlistResult:
    rows = parse_metrics_csv(csv_path)
    sessions = build_anchor_sessions(rows)
    asr_index = load_asr_directory(asr_dir)
    profiles = parse_potential_excel(excel_path)

    missing = []
    for session in sessions:
        session.sample_asr = resolve_asr(asr_index, session.serial_no, session.live_room_id)
        if not session.sample_asr:
            missing.append(f"{session.serial_no}_{session.live_room_id}")
        session.potential = get_profile(profiles, session.serial_no)

    if missing:
        print(
            f"警告: {len(missing)} 条未匹配 ASR: {', '.join(missing[:5])}"
            + (" ..." if len(missing) > 5 else ""),
            file=sys.stderr,
        )

    result = run_scoring(sessions, str(csv_path), str(asr_dir), str(excel_path))
    export_json(result, output_path)
    return result


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="潜力主播海选系统 — 单文件版")
    p.add_argument("-i", "--input", default=DEFAULT_CSV, help="指标 CSV")
    p.add_argument("-a", "--asr-dir", default=DEFAULT_ASR_DIR, help="ASR 目录")
    p.add_argument("-p", "--potential-excel", default=DEFAULT_EXCEL, help="潜能 Excel")
    p.add_argument("-o", "--output", default=DEFAULT_OUTPUT, help="输出 JSON")
    return p


def main(argv: Optional[List[str]] = None) -> int:
    args = build_arg_parser().parse_args(argv)
    try:
        result = run_pipeline(
            Path(args.input).resolve(),
            Path(args.asr_dir).resolve(),
            Path(args.potential_excel).resolve(),
            Path(args.output),
        )
    except (FileNotFoundError, ValueError, ImportError) as e:
        print(f"错误: {e}", file=sys.stderr)
        return 1

    hard_n = sum(1 for x in result.items if x.status == "硬性淘汰")
    soft_n = sum(1 for x in result.items if x.status == "综合淘汰")
    print(
        f"完成: 共 {result.total_input} 条, 入围 {result.total_passed}, "
        f"硬性淘汰 {hard_n}, 综合淘汰 {soft_n}, 输出 -> {args.output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
